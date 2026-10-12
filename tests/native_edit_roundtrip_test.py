#!/usr/bin/env python3
"""Native PPTX editing regression: change text and drag shapes, then round-trip."""
from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT=Path(__file__).resolve().parents[1]


def descendants(collection):
    for shape in collection:
        yield shape
        if hasattr(shape,"shapes"):
            yield from descendants(shape.shapes)


for version,pages in (("v7",44),("v10",55),("current",107)):
    path=ROOT/"downloads"/f"agent101-{version}-editable.pptx"
    pres=Presentation(path)
    assert len(pres.slides)==pages
    shapes=list(descendants(pres.slides[0].shapes))
    source=next(s for s in shapes if s.has_text_frame and s.text_frame.text.strip())
    old_left=source.left
    marker=f"AC58 - {version} EDITABLE OBJECT"
    source.text_frame.text=marker
    source.left=old_left+Inches(.25)

    # Cards/flowchart primitives must also be independently movable:
    # source picture-only presentations have no such native shapes.
    card=next(
        s for s in descendants(pres.slides[2].shapes)
        if s.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE
        and s.width < pres.slide_width*.80
        and s.height < pres.slide_height*.80
    )
    card_id=card.shape_id
    card_x=card.left
    card.left=card_x+Inches(.125)

    # An SVG workflow arrow is a separate movable DrawingML freeform.
    arrow_slide = 2 if version=="v7" else 1
    arrow_id=arrow_x=None
    arrow_page=None
    if version in ("v7","current"):
        arrow_page=pres.slides[arrow_slide]
        arrow=next(sh for sh in descendants(arrow_page.shapes)
                   if sh.shape_type==MSO_SHAPE_TYPE.FREEFORM)
        arrow_id=arrow.shape_id
        arrow_x=arrow.left
        arrow.left=arrow_x+Inches(.1)

    with TemporaryDirectory(prefix="agent101-native-edit-") as tmp:
        target=Path(tmp)/"modified.pptx"
        pres.save(target)
        reopened=Presentation(target)
        found=list(descendants(reopened.slides[0].shapes))
        assert any(s.has_text_frame and s.text_frame.text==marker
                   and s.left >= old_left+Inches(.24) for s in found),version
        found_cards=list(descendants(reopened.slides[2].shapes))
        assert any(s.shape_id==card_id and s.left>=card_x+Inches(.12)
                   for s in found_cards),(version,"native diagram/card move")
        if arrow_id is not None:
            arrow_new=list(descendants(reopened.slides[arrow_slide].shapes))
            assert any(sh.shape_id==arrow_id and sh.left>=arrow_x+Inches(.09)
                       for sh in arrow_new),(version,"native arrow move")
        assert len(reopened.slides)==pages
    print(f"{version}: {pages} slides, text + card/diagram drag + save/reopen PASS")
print("ALL_NATIVE_EDIT_ROUNDTRIPS_PASS")
