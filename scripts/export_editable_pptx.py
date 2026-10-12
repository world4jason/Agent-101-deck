#!/usr/bin/env python3
"""Export Agent 101 HTML versions as native, editable PowerPoint objects.

DOM/CSS -> primitive SVG -> PowerPoint DrawingML. Existing picture-only
exports remain intact as a clearly marked static fallback.

The editable SVG extractor and DrawingML converter are derived from
GX-Alex/html2pptx (MIT); see scripts/vendor/html2pptx/LICENSE.
"""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
VENDOR = ROOT / "scripts" / "vendor" / "html2pptx"
sys.path.insert(0, str(VENDOR))
from svg_to_pptx import create_pptx_with_native_svg  # noqa: E402

VERSIONS = {
    "v7": ("versions/v7/index.html", 44),
    "v10": ("drafts/rev10.html", 55),
    "current": ("slides/index.html", 107),
}
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

# The display website has only one active section at a time. This export-only
# stylesheet adjusts its scene to a fixed 16:9 canvas without touching the
# authoritative HTML, stylesheets, or history versions.
EXPORT_CSS = """
<style id="agent101-native-pptx-only">
html,body {width:1280px!important;height:720px!important;min-height:0!important;overflow:hidden!important;margin:0!important}
body .deck,body .slides,body .rev10-draft .deck {
  position:fixed!important;inset:0!important;transform:none!important;
  width:1280px!important;height:720px!important;overflow:hidden!important;
}
section.deck-slide {
  position:fixed!important;inset:0!important;transform:none!important;
  width:1280px!important;height:720px!important;min-height:0!important;
  max-width:none!important;opacity:1!important;visibility:visible!important;
  box-sizing:border-box!important;
}
/* PowerPoint has no browser reveal buttons: export final readable state.
   State-only .move-step markers still honor their original display:none. */
section.deck-slide .fragment:not(.move-step) {
  opacity:1!important;visibility:visible!important;transform:none!important;
}
.topbar,.controls,.toggle-controls,.deck-version-tools {display:none!important}
</style>
"""


def prepare_html(source: Path, dest: Path, max_pages: int | None) -> int:
    original = source.read_text(encoding="utf-8")
    selected = 0

    def mark(match: re.Match[str]) -> str:
        nonlocal selected
        selected += 1
        if max_pages is not None and selected > max_pages:
            return match.group(0)
        attrs = match.group(1)
        return '<section class="deck-slide slide' + attrs + '" data-export-display="flex"'

    # The existing classes start exactly with slide; matching only <section>
    # avoids accidentally changing nested cards or SVG elements.
    html = re.sub(r'<section class="slide([^"]*)"', mark, original)
    html = re.sub(r"</head>", EXPORT_CSS + "</head>", html, count=1, flags=re.I)
    if selected == 0 or "<section class=\"deck-slide" not in html:
        raise ValueError(f"Slide sections not found in {source}")
    dest.write_text(html, encoding="utf-8")
    return selected


def notes_from_existing_export(path: Path) -> dict[str, str]:
    if not path.exists():
        raise FileNotFoundError(
            f"Speaker-notes source missing: {path}. The static PPTX is required "
            "to preserve lecture notes and Exit Check answers."
        )
    presentation = Presentation(str(path))
    result = {}
    for index, slide in enumerate(presentation.slides, 1):
        original = slide.notes_slide.notes_text_frame.text
        original = original.replace(
            "此 PPTX 以圖片保留網頁投影片版面；動畫、逐步揭露與按鈕須使用網頁版本。",
            "此 PPTX 使用可編輯的 PowerPoint 原生文字與圖形；網頁動畫、逐步揭露與按鈕仍須使用網頁版本。"
        )
        result[f"{index:03d}_slide"] = original
    return result


def inspect_native_package(path: Path, expected_pages: int) -> dict:
    with ZipFile(path) as zf:
        slide_files = [
            name for name in zf.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        ]
        width_emu = height_emu = 0
        pres = ET.fromstring(zf.read("ppt/presentation.xml"))
        sld_size = pres.find("p:sldSz", NS)
        if sld_size is not None:
            width_emu = int(sld_size.get("cx", 0))
            height_emu = int(sld_size.get("cy", 0))
        if len(slide_files) != expected_pages:
            raise AssertionError(f"{path.name}: {len(slide_files)} pages != {expected_pages}")
        totals = {"slides": len(slide_files), "text_shapes": 0,
                  "other_shapes": 0, "pictures": 0, "text_chars": 0,
                  "slides_without_editable_text": [], "full_slide_pictures": []}
        for name in slide_files:
            index = int(re.search(r"slide(\d+)\.xml", name).group(1))
            xml = ET.fromstring(zf.read(name))
            shape_nodes = xml.findall(".//p:sp", NS)
            pic_nodes = xml.findall(".//p:pic", NS)
            text_shapes = 0
            for sp in shape_nodes:
                runs = sp.findall(".//a:t", NS)
                if any((r.text or "").strip() for r in runs):
                    text_shapes += 1
                    totals["text_chars"] += sum(len(r.text or "") for r in runs)
            totals["text_shapes"] += text_shapes
            totals["other_shapes"] += len(shape_nodes) - text_shapes
            totals["pictures"] += len(pic_nodes)
            if text_shapes == 0:
                totals["slides_without_editable_text"].append(index)
            for pic in pic_nodes:
                ext = pic.find("./p:spPr/a:xfrm/a:ext", NS)
                if ext is None or not width_emu or not height_emu:
                    continue
                area = int(ext.get("cx", 0)) * int(ext.get("cy", 0))
                if area >= width_emu * height_emu * 0.9:
                    totals["full_slide_pictures"].append(index)
        # Do not allow a screenshot covered by invisible text shapes to pass.
        if totals["full_slide_pictures"]:
            raise AssertionError(
                f"{path.name}: flattened images on {totals['full_slide_pictures']}"
            )
        # Some genuine screenshot-only slide can still have no editable text,
        # but every Agent-101 page has a heading, so failing is intentional.
        if totals["slides_without_editable_text"]:
            raise AssertionError(
                f"{path.name}: pages without editable text "
                f"{totals['slides_without_editable_text']}"
            )
        if totals["text_shapes"] < expected_pages * 2:
            raise AssertionError(
                f"{path.name}: only {totals['text_shapes']} text shapes "
                f"across {expected_pages} pages"
            )
        return totals




def text_recall_by_page(project: Path, output: Path, pages: int) -> list[dict]:
    """Compare browser-visible text with native PPTX text, without OCR.

    Uses normalized character trigrams, so nonsemantic wrapping and whitespace
    differences do not cause false failures. This is an audit signal, not a
    replacement for actually inspecting visual output.
    """
    from collections import Counter
    results = []
    with ZipFile(output) as zf:
        for index in range(1, pages + 1):
            html_visible = (project / "notes" / f"{index:03d}_slide.md").read_text(
                encoding="utf-8").split("\n\n", 1)[-1]
            ppt_root = ET.fromstring(zf.read(f"ppt/slides/slide{index}.xml"))
            native = " ".join(
                (node.text or "") for node in ppt_root.findall(".//a:t", NS)
            )
            compact = lambda t: re.sub(r"[\W_]+", "", t.lower())
            a, b = compact(html_visible), compact(native)
            ngrams = lambda t: Counter(t[i:i+3] for i in range(max(len(t)-2,0)))
            src, dst = ngrams(a), ngrams(b)
            denom = sum(src.values())
            hit = sum((src & dst).values())
            results.append({
                "page": index, "html_chars": len(a), "pptx_chars": len(b),
                "trigram_recall": round(hit / denom, 4) if denom else 1.0
            })
    return results


def export_version(version: str, count: int, max_pages: int | None,
                   output_dir: Path, retain_intermediate: bool) -> dict:
    rel, expected = VERSIONS[version]
    source = ROOT / rel
    requested = min(max_pages or expected, expected)
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"agent101-{version}-editable.pptx"

    # Temp source MUST sit beside original HTML so ../ CSS and image links work.
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".html", prefix=".agent101-native-export-",
        dir=source.parent, delete=False
    ) as handle:
        prepared = Path(handle.name)
    try:
        actual_sections = prepare_html(source, prepared, max_pages)
        if actual_sections != expected:
            raise AssertionError(
                f"{version}: source changed: {actual_sections} != {expected}"
            )
        with tempfile.TemporaryDirectory(prefix=f"agent101-{version}-native-") as tmp:
            project = Path(tmp) / "project"
            env = dict(os.environ)
            chrome = env.get("CHROME")
            if not chrome:
                for candidate in [
                    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
                    "/usr/bin/chromium", "/usr/bin/google-chrome"
                ]:
                    if Path(candidate).exists():
                        env["CHROME"] = candidate
                        break
            print(f"[native] {version}: extracting {requested}/{expected} editable pages", flush=True)
            subprocess.run([
                "node", str(VENDOR / "html_dom_to_editable_svg.js"),
                str(prepared), str(project)
            ], check=True, env=env)
            svg_files = sorted((project / "svg_output").glob("*_slide.svg"))
            if len(svg_files) != requested:
                raise AssertionError(
                    f"{version}: SVG output {len(svg_files)} != {requested}"
                )
            notes = notes_from_existing_export(
                ROOT / "downloads" / f"agent101-{version}.pptx"
            )
            notes = {f"{i:03d}_slide": notes.get(f"{i:03d}_slide", "")
                     for i in range(1, requested + 1)}
            ok = create_pptx_with_native_svg(
                svg_files=svg_files,
                output_path=output,
                canvas_format="ppt169",
                verbose=False,
                transition=None,
                use_compat_mode=False,
                notes=notes,
                enable_notes=True,
                use_native_shapes=True,
                animation=None,
            )
            if not ok or not output.exists():
                raise RuntimeError(f"Native DrawingML conversion failed: {version}")
            recall = text_recall_by_page(project, output, requested)
            worst = sorted(recall, key=lambda r:r["trigram_recall"])[:8]
            print(f"[native] {version} visible-text recall (worst): {worst}", flush=True)
            low = [r for r in recall if r["trigram_recall"] < .80]
            if low:
                raise AssertionError(
                    f"{version}: important visible text may be missing from PPTX "
                    f"(trigram recall < 80%): {low}"
                )
            if retain_intermediate:
                import shutil
                debug = output_dir / ("native-debug-" + version)
                if debug.exists():
                    shutil.rmtree(debug)
                shutil.copytree(project, debug, ignore=shutil.ignore_patterns("chrome-profile"))
        stats = inspect_native_package(output, requested)
        stats["text_recall"] = {
            "minimum": min(x["trigram_recall"] for x in recall),
            "mean": round(sum(x["trigram_recall"] for x in recall)/len(recall),4),
            "worst_pages": sorted(recall, key=lambda r:r["trigram_recall"])[:8]
        }
        stats["version"] = version
        stats["file"] = output.name
        stats["bytes"] = output.stat().st_size
        print(f"[native] PASS {version}: {stats}", flush=True)
        return stats
    finally:
        prepared.unlink(missing_ok=True)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--version", choices=["all", *VERSIONS], default="all")
    ap.add_argument("--max-pages", type=int, default=None)
    ap.add_argument("--output-dir", type=Path, default=ROOT / "downloads")
    ap.add_argument("--keep-intermediate", action="store_true")
    args = ap.parse_args()
    if args.max_pages is not None and args.max_pages <= 0:
        ap.error("--max-pages must be a positive integer")
    versions = list(VERSIONS) if args.version == "all" else [args.version]
    records = [export_version(version, VERSIONS[version][1],
                              args.max_pages, args.output_dir,
                              args.keep_intermediate)
               for version in versions]
    report = args.output_dir / "agent101-native-export-report.json"
    report.write_text(json.dumps({
        "pipeline": "HTML/CSS/SVG -> editable DrawingML, not screenshots",
        "results": records,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"EDITABLE_PPTX_OK {report}", flush=True)


if __name__ == "__main__":
    main()
